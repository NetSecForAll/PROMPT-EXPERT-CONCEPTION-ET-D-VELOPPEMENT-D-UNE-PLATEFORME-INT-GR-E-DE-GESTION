import pytest

from apps.journal.models import EvenementSecurite

PAGES_ADMINISTRATION = [
    "/administration/sites/",
    "/administration/services/",
    "/administration/agents/",
    "/administration/comptes/",
    "/administration/parametres/",
    "/validations/circuits/",
    "/journal/",
]


def test_anonyme_redirige_vers_connexion(client, db):
    reponse = client.get("/")
    assert reponse.status_code == 302 and reponse.url.startswith("/connexion/")


@pytest.mark.parametrize("url", PAGES_ADMINISTRATION)
def test_agent_sans_droit_recoit_403_journalise(equipe, connecter, url):
    client = connecter(equipe["agent"])
    reponse = client.get(url)
    assert reponse.status_code == 403
    assert EvenementSecurite.objects.filter(type="acces_refuse", utilisateur=equipe["agent"], chemin=url).exists()


def test_export_excel_refuse_sans_droit(equipe, connecter):
    client = connecter(equipe["agent"])
    assert client.get("/administration/agents/?export=xlsx").status_code == 403


@pytest.mark.parametrize("url", PAGES_ADMINISTRATION)
def test_admin_fonctionnel_accede(equipe, connecter, url):
    assert connecter(equipe["admin_fonc"]).get(url).status_code == 200


def test_auditeur_lit_mais_ne_modifie_pas(equipe, connecter, organisation):
    client = connecter(equipe["auditeur"])
    assert client.get("/administration/agents/").status_code == 200
    assert client.get("/journal/").status_code == 200
    assert client.get("/administration/agents/nouveau/").status_code == 403
    assert (
        client.post(f"/administration/sites/{organisation['site'].pk}/", {"code": "X", "nom": "X"}).status_code == 403
    )


def test_admin_technique_sans_acces_journal_ni_referentiels(equipe, connecter):
    client = connecter(equipe["admin_tech"])
    assert client.get("/administration/comptes/").status_code == 200
    assert client.get("/journal/").status_code == 403
    assert client.get("/administration/agents/nouveau/").status_code == 403


def test_menu_selon_les_droits(equipe, connecter):
    contenu_agent = connecter(equipe["agent"]).get("/").content.decode()
    assert "/administration/comptes/" not in contenu_agent
    contenu_admin = connecter(equipe["admin_fonc"]).get("/").content.decode()
    assert "/administration/comptes/" in contenu_admin


def test_export_excel(equipe, connecter):
    from io import BytesIO

    from openpyxl import load_workbook

    reponse = connecter(equipe["admin_fonc"]).get("/administration/agents/?export=xlsx&service=")
    assert reponse.status_code == 200
    feuille = load_workbook(BytesIO(reponse.content)).active
    assert feuille.cell(1, 1).value == "Matricule"
    assert feuille.max_row == 1 + 8  # en-tête + 8 agents de l'équipe
    assert EvenementSecurite.objects.filter(type="export").exists()


def test_recherche_tri_pagination(equipe, connecter):
    client = connecter(equipe["admin_fonc"])
    contenu = client.get("/administration/agents/?q=auditeur").content.decode()
    assert "Auditeur" in contenu and "Directeur" not in contenu
    assert client.get("/administration/agents/?tri=-nom").status_code == 200
    # Un tri non autorisé est ignoré, pas exécuté.
    assert client.get("/administration/agents/?tri=utilisateur__password").status_code == 200


def test_en_tetes_de_securite(equipe, connecter):
    reponse = connecter(equipe["agent"]).get("/")
    assert "script-src 'self'" in reponse["Content-Security-Policy"]
    assert reponse["X-Frame-Options"] == "DENY"


def test_bouton_ouvrir_masque_sans_droit_de_modification(equipe, connecter):
    lien = f'href="/administration/agents/{equipe["agent"].agent_id}/"'
    assert lien not in connecter(equipe["auditeur"]).get("/administration/agents/").content.decode()
    assert lien in connecter(equipe["admin_fonc"]).get("/administration/agents/").content.decode()
